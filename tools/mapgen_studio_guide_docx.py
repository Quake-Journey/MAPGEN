r"""The Studio's guide as DOCX beside its Markdown (Fable's brief 10, D1; the PO, 06.10: «я буду смотреть DOCX»).

    python tools/mapgen_studio_guide_docx.py

doc/mapgen_studio/MAPGEN_Studio_Guide_RU.md and _EN.md -> .docx next to them, every picture EMBEDDED (no linked file,
no field to update - `memory/feedback_docx_builds_must_open_cleanly.md`), then checked by
`tools/validate_q2prox_docx.py` when it is there. Only what the guides use: headings, paragraphs, one-level lists with
continued lines, **bold**, `code`, ![alt](picture) on a line of its own.
"""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

TOOLS = Path(__file__).resolve().parent
DOC = TOOLS.parent / "doc" / "mapgen_studio"
VALIDATOR = TOOLS / "validate_q2prox_docx.py"


def runs(par, text: str) -> None:
    """**bold** and `code` as runs of the paragraph."""
    for piece in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
        else:
            par.add_run(piece)


def build(md: Path) -> Path:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Segoe UI"
    normal.font.size = Pt(10.5)
    section = doc.sections[0]
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(1.8)
    width = section.page_width - section.left_margin - section.right_margin
    para, item = [], None

    def flush():
        nonlocal item
        if para:
            runs(doc.add_paragraph(), " ".join(para))
            para.clear()
        item = None

    for line in md.read_text(encoding="utf-8").splitlines():
        s = line.rstrip()
        img = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", s.strip())
        if not s.strip():
            flush()
        elif img:
            flush()
            pic = (md.parent / img.group(2)).resolve()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(pic), width=width)
            cap = doc.add_paragraph(img.group(1))
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].italic = True
            cap.runs[0].font.size = Pt(9)
        elif s.startswith("#"):
            flush()
            level = len(s) - len(s.lstrip("#"))
            doc.add_heading(s[level:].strip(), level=0 if level == 1 else level - 1)
        elif s.startswith("- "):
            if para:
                flush()
            item = doc.add_paragraph(style="List Bullet")
            item._text = s[2:]
            runs(item, s[2:])
        elif item is not None and s.startswith("  "):
            # a continued list line: rebuild the item with the whole text
            text = item._text + " " + s.strip()
            for r in list(item.runs):
                r._element.getparent().remove(r._element)
            item._text = text
            runs(item, text)
        else:
            item = None
            para.append(s.strip())
    flush()
    out = md.with_suffix(".docx")
    doc.save(out)
    links_up_to_date(out)
    return out


def links_up_to_date(path: Path) -> None:
    """python-docx's own template says LinksUpToDate false; Word then may ask about updating links. Nothing here is
    linked - every picture is embedded - so the package says so."""
    import shutil
    import zipfile
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "docProps/app.xml":
                data = data.replace(b"<LinksUpToDate>false</LinksUpToDate>", b"<LinksUpToDate>true</LinksUpToDate>")
            dst.writestr(item, data)
    shutil.move(tmp, path)


def main() -> int:
    bad = 0
    for lang in ("RU", "EN"):
        out = build(DOC / f"MAPGEN_Studio_Guide_{lang}.md")
        print(out)
        if VALIDATOR.is_file():
            run = subprocess.run([sys.executable, str(VALIDATOR), str(out)], capture_output=True, text=True)
            print((run.stdout + run.stderr).strip()[-600:])
            bad += run.returncode != 0
    return 1 if bad else 0


if __name__ == "__main__":
    raise SystemExit(main())
